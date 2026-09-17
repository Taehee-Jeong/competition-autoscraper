# -*- coding: utf-8 -*-
"""엑셀 출력: '확정 컴피티션 공지 목록'과 동일한 컬럼 구조의 후보 목록.

- data/후보_컴피티션_목록.xlsx 에 누적 저장
- 공식 링크(또는 위비티 ix) 기준 중복 제거 → 새 대회만 추가
- 새로 추가된 건은 '공지 상태' = '신규후보'
"""
import os
import re
import logging
from datetime import date
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment

log = logging.getLogger("output")

COLUMNS = ["공지월", "학부", "대회명", "주최 / 주관", "분야 / 주제", "신청 마감일",
           "대회 일정", "팀 인원 (개인 참가 여부)", "참가 자격", "공식 링크",
           "상금 / 혜택", "참가비", "공지 상태", "담당자", "비고"]

# 검토상태 → 엑셀 '공지 상태' 표기
STATUS_LABEL = {"자동통과": "신규후보",
                "확인필요": "신규후보(확인필요)",
                "LLM판정_부적합": "신규후보(LLM의심)"}

HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(name="Arial", bold=True, color="FFFFFF")
BODY_FONT = Font(name="Arial", size=10)


def _cell(v):
    """엑셀이 받아주는 값으로 낮춘다.

    한 칸의 타입 하나 때문에 수집 전체가 저장 직전에 죽는 일이 실제로 있었다
    (LLM이 상금을 dict로 반환 → openpyxl ValueError → 26분치 손실).
    """
    return v if v is None or isinstance(v, (str, int, float)) else str(v)


def _row_from_item(it: dict) -> list:
    deadline = it.get("deadline") or "확인필요"
    notice_month = date.today().strftime("%Y-%m")
    link = it.get("homepage") or it.get("url", "")
    eligibility = it.get("eligibility_llm") or it.get("target") or "확인필요"
    src_name = it.get("source", "위비티")
    loc = f", 장소 {it['location']}" if it.get("location") else ""
    remark = f"{src_name} 수집 / 판정: 재학생 {it.get('판정_재학생','?')}, " \
             f"팀 {it.get('판정_팀','?')}, D-{it.get('days_left','?')}{loc}"
    return [_cell(v) for v in
            [notice_month, it.get("학부", "미분류"), it.get("title", ""),
             it.get("host", ""), it.get("field", ""), deadline,
             it.get("schedule", "확인필요"), it.get("team_size", "확인필요"),
             eligibility, link, it.get("prize", ""), it.get("fee", "확인필요"),
             STATUS_LABEL.get(it.get("검토상태"), "신규후보(확인필요)"),
             "", remark]]


def _dedup_key(text) -> str:
    """중복 판정용 대회명 키. 공백·괄호·문장부호를 지우고 소문자로 맞춘다.

    같은 대회가 사이트마다 '대학가요제'/'대학 가요제', '『…』'/'…' 처럼 표기만
    다르게 올라오기 때문이다. 반대로 단어가 더 붙은 경우('…학술대회' vs
    '…학술대회 공모전')는 일부러 잡지 않는다 — 유사도 매칭까지 가면 서로 다른
    대회('Global Hack Week: Data' vs '… : Agents')를 합쳐버린다.
    """
    return re.sub(r"[^0-9a-z가-힣]", "", str(text or "").lower())


def _existing_keys(path: str) -> set[str]:
    """기존 시트에서 중복 판정 키(공식 링크 + 정규화한 대회명)를 모은다."""
    if not os.path.exists(path):
        return set()
    wb = load_workbook(path, read_only=True)
    ws = wb.active
    keys = set()
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row:
            continue
        if len(row) >= 10 and row[9]:
            keys.add(str(row[9]).strip())
        if len(row) >= 3 and row[2]:
            keys.add(_dedup_key(row[2]))  # 대회명으로도 중복 방지
    keys.discard("")  # 빈 값이 들어가면 링크 없는 항목이 서로를 막는다
    return keys


def write_candidates(items: list[dict], path: str) -> int:
    """새 후보를 파일에 추가. 반환값: 추가된 건수."""
    known = _existing_keys(path)

    if os.path.exists(path):
        wb = load_workbook(path)
        ws = wb.active
    else:
        wb = Workbook()
        ws = wb.active
        ws.title = "후보목록"
        ws.append(COLUMNS)
        for c, w in zip(ws[1], [10, 26, 40, 24, 20, 12, 26, 22, 26, 40, 22, 10, 14, 10, 34]):
            c.fill, c.font = HEADER_FILL, HEADER_FONT
            c.alignment = Alignment(horizontal="center", vertical="center")
            ws.column_dimensions[c.column_letter].width = w
        ws.freeze_panes = "A2"

    added = 0
    for it in items:
        link = (it.get("homepage") or it.get("url", "")).strip()
        title_key = _dedup_key(it.get("title"))
        if (link and link in known) or (title_key and title_key in known):
            continue
        ws.append(_row_from_item(it))
        for cell in ws[ws.max_row]:
            cell.font = BODY_FONT
            cell.alignment = Alignment(vertical="top", wrap_text=True)
        known.update(k for k in (link, title_key) if k)
        added += 1

    wb.save(path)
    log.info("엑셀 저장: %s (신규 %d건)", path, added)
    return added
