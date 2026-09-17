# -*- coding: utf-8 -*-
"""검수용 SQLite 저장소.

엑셀은 '공지용 산출물'로 남기고, 사람이 판단한 결과는 여기에 쌓는다.
엑셀 셀에 직접 쓰면 '확정'만 읽을 수 있고 '왜 아닌지'는 남지 않아서,
진화 엔진이 부정 사례를 영영 못 받는다 (DEVELOPMENT.md §9-1).
"""
import os
import json
import sqlite3
from datetime import date, datetime

BASE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DB_PATH = os.environ.get("WEB_DB", os.path.join(BASE, "data", "review.sqlite3"))

# 기각 사유 → 진화 엔진이 쓸 라벨. None이면 학습 데이터로 쓰지 않는다.
REJECT_REASONS = {
    "재학생불가": {"label": "재학생이 참가할 수 없음", "teaches": ("student_ok", False)},
    "팀불가":    {"label": "팀 참가가 안 됨",       "teaches": ("team_ok", False)},
    "마감임박":   {"label": "마감이 너무 가까움",     "teaches": None},
    "중복":      {"label": "다른 줄과 같은 대회",    "teaches": None},
    "분야벗어남": {"label": "우리 분야가 아님",       "teaches": None},
    "기타":      {"label": "기타",                "teaches": None},
}
DECISIONS = ("확정", "기각", "보류")

# 일괄 확정 기본 기준. 규칙 3종이 다 통과(+50)하고 마감 여유가 있어야 닿는 값이다.
AUTO_THRESHOLD = 75

SCHEMA = """
CREATE TABLE IF NOT EXISTS candidates (
    id           INTEGER PRIMARY KEY,
    dedup_key    TEXT    NOT NULL UNIQUE,
    title        TEXT    NOT NULL,
    link         TEXT,
    source       TEXT,
    host         TEXT,
    field        TEXT,
    school       TEXT,
    deadline     TEXT,
    schedule     TEXT,
    team_size    TEXT,
    eligibility  TEXT,
    prize        TEXT,
    fee          TEXT,
    body         TEXT,
    auto_status  TEXT,
    judge_student TEXT,
    judge_team    TEXT,
    judge_deadline TEXT,
    remark       TEXT,
    first_seen   TEXT NOT NULL,
    last_seen    TEXT NOT NULL
);
CREATE TABLE IF NOT EXISTS reviews (
    candidate_id INTEGER PRIMARY KEY REFERENCES candidates(id) ON DELETE CASCADE,
    decision     TEXT NOT NULL,
    reason       TEXT,
    owner        TEXT,
    note         TEXT,
    decided_at   TEXT NOT NULL
);
CREATE TABLE IF NOT EXISTS ai_reviews (
    candidate_id INTEGER PRIMARY KEY REFERENCES candidates(id) ON DELETE CASCADE,
    decision     TEXT NOT NULL,
    reason       TEXT,
    confidence   TEXT NOT NULL,
    rationale    TEXT,
    judged_at    TEXT NOT NULL
);
CREATE TABLE IF NOT EXISTS merges (
    dup_id     INTEGER PRIMARY KEY REFERENCES candidates(id) ON DELETE CASCADE,
    keep_id    INTEGER NOT NULL REFERENCES candidates(id) ON DELETE CASCADE,
    merged_at  TEXT NOT NULL
);
CREATE TABLE IF NOT EXISTS runs (
    id          INTEGER PRIMARY KEY,
    started_at  TEXT NOT NULL,
    finished_at TEXT,
    status      TEXT NOT NULL,
    collected   INTEGER,
    passed      INTEGER,
    added       INTEGER,
    log_path    TEXT
);
CREATE INDEX IF NOT EXISTS idx_cand_deadline ON candidates(deadline);
CREATE INDEX IF NOT EXISTS idx_cand_school   ON candidates(school);
"""


def connect() -> sqlite3.Connection:
    os.makedirs(os.path.dirname(DB_PATH), exist_ok=True)
    conn = sqlite3.connect(DB_PATH, timeout=15)
    conn.row_factory = sqlite3.Row
    conn.execute("PRAGMA foreign_keys = ON")
    return conn


def _add_column(conn, table: str, col: str, decl: str):
    if col not in {r[1] for r in conn.execute(f"PRAGMA table_info({table})")}:
        conn.execute(f"ALTER TABLE {table} ADD COLUMN {col} {decl}")


def init():
    with connect() as conn:
        conn.executescript(SCHEMA)
        # 사람 판단인지 AI 제안을 적용한 것인지 구분한다. 이걸 안 나누면 AI가 만든
        # 라벨이 평가셋에 섞여 들어가 모델이 제 출력으로 제 규칙을 학습하게 된다.
        _add_column(conn, "reviews", "source", "TEXT NOT NULL DEFAULT 'human'")
        _add_column(conn, "runs", "kind", "TEXT NOT NULL DEFAULT 'scrape'")


# ------------------------------------------------------------------ 적재
def upsert_candidates(items: list[dict]) -> tuple[int, int]:
    """파이프라인 결과를 반영. → (신규, 갱신) 건수.

    이미 있는 대회는 검수 결과를 유지한 채 마감일·본문 같은 값만 새로 고친다.
    """
    from scraper.output import _dedup_key
    now = datetime.now().isoformat(timespec="seconds")
    new = upd = 0
    with connect() as conn:
        for it in items:
            key = _dedup_key(it.get("title"))
            if not key:
                continue
            row = {
                "dedup_key": key,
                "title": (it.get("title") or "").strip(),
                "link": (it.get("homepage") or it.get("url") or "").strip(),
                "source": it.get("source", ""),
                "host": it.get("host", ""),
                "field": it.get("field", ""),
                "school": it.get("학부", "미분류"),
                "deadline": it.get("deadline"),
                "schedule": it.get("schedule", "확인필요"),
                "team_size": it.get("team_size", "확인필요"),
                "eligibility": it.get("eligibility_llm") or it.get("target") or "확인필요",
                "prize": it.get("prize", ""),
                "fee": it.get("fee", "확인필요"),
                "body": (it.get("body") or "")[:20000],
                "auto_status": it.get("검토상태", "확인필요"),
                "judge_student": it.get("판정_재학생", ""),
                "judge_team": it.get("판정_팀", ""),
                "judge_deadline": it.get("판정_마감", ""),
                "remark": f"D-{it.get('days_left', '?')}",
                "last_seen": now,
            }
            cur = conn.execute("SELECT id FROM candidates WHERE dedup_key = ?", (key,))
            hit = cur.fetchone()
            if hit:
                cols = ", ".join(f"{k} = :{k}" for k in row if k != "dedup_key")
                conn.execute(f"UPDATE candidates SET {cols} WHERE dedup_key = :dedup_key", row)
                upd += 1
            else:
                row["first_seen"] = now
                cols = ", ".join(row)
                conn.execute(f"INSERT INTO candidates ({cols}) "
                             f"VALUES ({', '.join(':' + k for k in row)})", row)
                new += 1
    return new, upd


def import_from_xlsx(path: str) -> int:
    """기존 엑셀 행을 DB로 옮긴다 (본문은 없으므로 검수만 가능). → 신규 건수."""
    from openpyxl import load_workbook
    from scraper.output import _dedup_key
    if not os.path.exists(path):
        return 0
    now = datetime.now().isoformat(timespec="seconds")
    ws = load_workbook(path, read_only=True).active
    added = 0
    with connect() as conn:
        for r in ws.iter_rows(min_row=2, values_only=True):
            if not r or len(r) < 15 or not r[2]:
                continue
            key = _dedup_key(r[2])
            if not key or conn.execute("SELECT 1 FROM candidates WHERE dedup_key = ?",
                                       (key,)).fetchone():
                continue
            conn.execute(
                "INSERT INTO candidates (dedup_key,title,link,source,host,field,school,"
                "deadline,schedule,team_size,eligibility,prize,fee,body,auto_status,"
                "judge_student,judge_team,judge_deadline,remark,first_seen,last_seen) "
                "VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,'',?,'','','',?,?,?)",
                (key, str(r[2]), str(r[9] or ""), str(r[14] or "").split(" 수집")[0],
                 str(r[3] or ""), str(r[4] or ""), str(r[1] or "미분류"),
                 str(r[5] or "") if r[5] and str(r[5]) != "확인필요" else None,
                 str(r[6] or ""), str(r[7] or ""), str(r[8] or ""), str(r[10] or ""),
                 str(r[11] or ""), str(r[12] or ""), str(r[14] or ""), now, now))
            added += 1
    return added


# ------------------------------------------------------------------ 조회
# 화면에 보이는 이름 → 정렬 방식. score/prize는 SQL로 못 하니 파이썬에서 다시 정렬한다.
SORT_LABELS = [("score", "적합도 높은순"), ("score_low", "적합도 낮은순"),
               ("deadline", "마감 임박순"), ("deadline_far", "마감 여유순"),
               ("prize", "상금 많은순"), ("school", "학부순"),
               ("source", "소스순"), ("recent", "최근 수집순"), ("title", "이름순")]
PY_SORTS = ("score", "score_low", "prize")

SORTS = {"deadline": "c.deadline IS NULL, c.deadline",
         "deadline_far": "c.deadline IS NULL, c.deadline DESC",
         "school": "c.school, c.deadline IS NULL, c.deadline",
         "source": "c.source, c.deadline IS NULL, c.deadline",
         "title": "c.title",
         "recent": "c.first_seen DESC",
         "score": "c.deadline IS NULL, c.deadline",
         "score_low": "c.deadline IS NULL, c.deadline",
         "prize": "c.deadline IS NULL, c.deadline"}


_UNIT = {"억": 100_000_000, "천만": 10_000_000, "백만": 1_000_000,
         "만": 10_000, "천": 1_000}


def prize_won(text) -> int:
    """상금 문구에서 대표 금액(원)을 뽑는다. 못 읽으면 0.

    표기가 제각각이다 — '3천만원~1천만원', '500만원', '$25,000',
    'Winner: INR 20,000', 데이콘은 그냥 '2000'(만원 단위).
    정확한 액수가 아니라 '큰 대회부터 보기' 위한 값이므로 최대값만 잡으면 충분하다.
    """
    import re
    s = str(text or "")
    if not s or s == "확인필요":
        return 0
    best = 0
    for num, unit in re.findall(r"([\d,.]+)\s*(억|천만|백만|만|천)?\s*원", s):
        try:
            best = max(best, int(float(num.replace(",", "")) * _UNIT.get(unit, 1)))
        except ValueError:
            pass
    if best:
        return best
    m = re.search(r"[$₩]\s*([\d,]+)", s)          # $25,000 → 대략 원화로
    if m:
        return int(m.group(1).replace(",", "")) * 1400
    m = re.search(r"^\s*([\d,]+)\s*$", s)         # 데이콘: 숫자만, 만원 단위
    if m:
        return int(m.group(1).replace(",", "")) * 10_000
    return 0


def list_candidates(status="", school="", source="", decision="", q="",
                    sort="deadline", hide_merged=True) -> list[sqlite3.Row]:
    where, args = ["1=1"], []
    if status:
        where.append("c.auto_status = ?"); args.append(status)
    else:
        where.append("c.auto_status != '탈락'")  # 기본: 탈락 제외
    if school:
        where.append("c.school = ?"); args.append(school)
    if source:
        where.append("c.source = ?"); args.append(source)
    if decision == "미검수":
        where.append("r.decision IS NULL")
    elif decision == "AI제안":
        where.append("r.decision IS NULL AND a.decision IS NOT NULL AND a.decision != '판단보류'")
    elif decision:
        where.append("r.decision = ?"); args.append(decision)
    if q:
        where.append("(c.title LIKE ? OR c.host LIKE ? OR c.field LIKE ?)")
        args += [f"%{q}%"] * 3
    if hide_merged:
        where.append("c.id NOT IN (SELECT dup_id FROM merges)")
    sql = (f"SELECT {SELECT_COLS} FROM candidates c {JOINS} "
           f"WHERE {' AND '.join(where)} ORDER BY {SORTS.get(sort, SORTS['deadline'])}")
    with connect() as conn:
        return conn.execute(sql, args).fetchall()


SELECT_COLS = ("c.*, r.decision, r.reason, r.owner, r.note, r.source AS review_source, "
               "a.decision AS ai_decision, a.reason AS ai_reason, "
               "a.confidence AS ai_confidence, a.rationale AS ai_rationale")
JOINS = ("LEFT JOIN reviews r ON r.candidate_id = c.id "
         "LEFT JOIN ai_reviews a ON a.candidate_id = c.id")


def get(cid: int):
    with connect() as conn:
        return conn.execute(f"SELECT {SELECT_COLS} FROM candidates c {JOINS} "
                            f"WHERE c.id = ?", (cid,)).fetchone()


def distinct(col: str) -> list[str]:
    if col not in ("school", "source", "auto_status"):
        raise ValueError(col)
    with connect() as conn:
        return [r[0] for r in conn.execute(
            f"SELECT DISTINCT {col} FROM candidates WHERE {col} != '' ORDER BY 1")]


def counts() -> dict:
    with connect() as conn:
        total = conn.execute("SELECT COUNT(*) FROM candidates").fetchone()[0]
        merged = conn.execute("SELECT COUNT(*) FROM merges").fetchone()[0]
        by = dict(conn.execute("SELECT decision, COUNT(*) FROM reviews GROUP BY 1").fetchall())
    by["미검수"] = total - merged - sum(by.values())
    by["전체"] = total - merged
    return by


def days_left(deadline: str | None) -> int | None:
    if not deadline:
        return None
    try:
        y, m, d = (int(x) for x in deadline.split("-")[:3])
        return (date(y, m, d) - date.today()).days
    except ValueError:
        return None


# ------------------------------------------------------------------ 검수
def set_review(cid: int, decision: str, reason="", owner="", note="", source="human"):
    if decision not in DECISIONS:
        raise ValueError(f"알 수 없는 판단: {decision}")
    if decision == "기각" and reason not in REJECT_REASONS:
        raise ValueError(f"알 수 없는 기각 사유: {reason}")
    with connect() as conn:
        conn.execute(
            "INSERT INTO reviews (candidate_id, decision, reason, owner, note, decided_at, source) "
            "VALUES (?,?,?,?,?,?,?) ON CONFLICT(candidate_id) DO UPDATE SET "
            "decision=excluded.decision, reason=excluded.reason, owner=excluded.owner, "
            "note=excluded.note, decided_at=excluded.decided_at, source=excluded.source",
            (cid, decision, reason if decision == "기각" else "", owner, note,
             datetime.now().isoformat(timespec="seconds"), source))


# ------------------------------------------------------------------ AI 검수
AI_CONFIDENCE = ("높음", "보통", "낮음")


def candidates_for_ai(redo: bool = False) -> list[sqlite3.Row]:
    """AI가 판정할 대상. 사람이 이미 판단한 건과 본문 없는 건은 제외한다."""
    sql = (f"SELECT {SELECT_COLS} FROM candidates c {JOINS} "
           "WHERE r.decision IS NULL AND c.body != '' "
           "AND c.id NOT IN (SELECT dup_id FROM merges) ")
    if not redo:
        sql += "AND a.decision IS NULL "
    with connect() as conn:
        return conn.execute(sql + "ORDER BY c.deadline IS NULL, c.deadline").fetchall()


def set_ai_review(cid: int, decision: str, reason="", confidence="낮음", rationale=""):
    if decision not in DECISIONS + ("판단보류",):
        raise ValueError(f"알 수 없는 AI 판단: {decision}")
    if confidence not in AI_CONFIDENCE:
        confidence = "낮음"
    with connect() as conn:
        conn.execute(
            "INSERT INTO ai_reviews (candidate_id, decision, reason, confidence, rationale, judged_at) "
            "VALUES (?,?,?,?,?,?) ON CONFLICT(candidate_id) DO UPDATE SET "
            "decision=excluded.decision, reason=excluded.reason, confidence=excluded.confidence, "
            "rationale=excluded.rationale, judged_at=excluded.judged_at",
            (cid, decision, reason if reason in REJECT_REASONS else "", confidence,
             rationale[:500], datetime.now().isoformat(timespec="seconds")))


def apply_ai_reviews(only: str = "", min_confidence: str = "") -> int:
    """AI 제안을 사람 판단 칸으로 옮긴다. → 적용 건수.

    only='확정'이면 확정 제안만, min_confidence='높음'이면 확신 높은 것만.
    source='ai'로 남기므로 평가셋에는 들어가지 않는다 (자기 출력으로 자기 학습 방지).
    """
    rows = [r for r in _pending_ai()
            if (not only or r["ai_decision"] == only)
            and (not min_confidence or r["ai_confidence"] == min_confidence)]
    for r in rows:
        set_review(r["id"], r["ai_decision"], r["ai_reason"] or "",
                   owner="AI", note="", source="ai")
    return len(rows)


def _pending_ai() -> list[sqlite3.Row]:
    with connect() as conn:
        return conn.execute(
            f"SELECT {SELECT_COLS} FROM candidates c {JOINS} "
            "WHERE r.decision IS NULL AND a.decision IN ('확정','기각') "
            "AND c.id NOT IN (SELECT dup_id FROM merges)").fetchall()


def ai_counts() -> dict:
    with connect() as conn:
        rows = conn.execute(
            "SELECT a.decision, a.confidence, COUNT(*) FROM ai_reviews a "
            "LEFT JOIN reviews r ON r.candidate_id = a.candidate_id "
            "WHERE r.decision IS NULL GROUP BY 1,2").fetchall()
    out = {"대기": 0, "확정": 0, "기각": 0, "판단보류": 0, "확신높음": 0}
    for dec, conf, n in rows:
        out[dec] = out.get(dec, 0) + n
        if dec in ("확정", "기각"):
            out["대기"] += n
            if conf == "높음":
                out["확신높음"] += n
    return out


def clear_review(cid: int):
    with connect() as conn:
        conn.execute("DELETE FROM reviews WHERE candidate_id = ?", (cid,))


def merge(dup_id: int, keep_id: int):
    if dup_id == keep_id:
        raise ValueError("자기 자신과는 합칠 수 없습니다")
    with connect() as conn:
        conn.execute("INSERT OR REPLACE INTO merges (dup_id, keep_id, merged_at) VALUES (?,?,?)",
                     (dup_id, keep_id, datetime.now().isoformat(timespec="seconds")))


def unmerge(dup_id: int):
    with connect() as conn:
        conn.execute("DELETE FROM merges WHERE dup_id = ?", (dup_id,))


# ------------------------------------------------------------------ 실행 기록
def start_run(log_path: str, kind: str = "scrape") -> int:
    with connect() as conn:
        cur = conn.execute(
            "INSERT INTO runs (started_at, status, log_path, kind) VALUES (?,?,?,?)",
            (datetime.now().isoformat(timespec="seconds"), "running", log_path, kind))
        return cur.lastrowid


def finish_run(run_id: int, status: str, collected=None, passed=None, added=None):
    with connect() as conn:
        conn.execute("UPDATE runs SET finished_at=?, status=?, collected=?, passed=?, added=? "
                     "WHERE id=?", (datetime.now().isoformat(timespec="seconds"), status,
                                    collected, passed, added, run_id))


def recent_runs(n=20) -> list[sqlite3.Row]:
    with connect() as conn:
        return conn.execute("SELECT * FROM runs ORDER BY id DESC LIMIT ?", (n,)).fetchall()


def running_run():
    with connect() as conn:
        return conn.execute("SELECT * FROM runs WHERE status='running' "
                            "ORDER BY id DESC LIMIT 1").fetchone()


# ------------------------------------------------------------------ 평가셋 내보내기
def golden_rows() -> list[dict]:
    """검수 결과 → 진화 엔진 평가셋. 확정은 긍정, 기각 사유는 부정 라벨이 된다."""
    rows = []
    with connect() as conn:
        for c in conn.execute(
                "SELECT c.*, r.decision, r.reason FROM candidates c "
                "JOIN reviews r ON r.candidate_id = c.id "
                "WHERE r.decision IN ('확정','기각') AND c.body != '' "
                "AND r.source = 'human'"):   # AI가 붙인 라벨로 규칙을 학습하면 안 된다
            rec = {"title": c["title"], "field": c["field"] or "",
                   "target": c["eligibility"] or "", "body": c["body"],
                   "lang": "ko" if any("가" <= ch <= "힣" for ch in c["title"]) else "en"}
            if c["decision"] == "확정":
                rec.update(student_ok=True, team_ok=True,
                           school=c["school"] if c["school"] != "미분류" else None)
            else:
                teaches = REJECT_REASONS.get(c["reason"], {}).get("teaches")
                if not teaches:
                    continue          # 마감·중복·분야는 규칙이 배울 게 없다
                rec[teaches[0]] = teaches[1]
            rows.append(rec)
    return rows


def write_goldenset(path: str) -> int:
    """평가셋.jsonl에 검수 결과를 덮어쓴다. 시드 12건은 별도 파일이므로 건드리지 않는다."""
    rows = golden_rows()
    with open(path, "w", encoding="utf-8") as f:
        for r in rows:
            f.write(json.dumps(r, ensure_ascii=False) + "\n")
    return len(rows)
