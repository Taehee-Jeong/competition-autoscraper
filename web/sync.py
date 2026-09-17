# -*- coding: utf-8 -*-
"""DB ↔ 엑셀.

수집 결과는 `db.upsert_candidates()`로 들어오고, 공지용 엑셀은 여기서 만든다.
엑셀은 더 이상 사람이 고치는 대상이 아니라 DB에서 뽑아내는 산출물이다.
"""
import os
import sys
from datetime import date

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from openpyxl import Workbook
from openpyxl.styles import Alignment

from scraper.output import (COLUMNS, STATUS_LABEL, HEADER_FILL,
                            HEADER_FONT, BODY_FONT)
from web import db, scoring

# 기존 15컬럼은 확정 공지 목록과 1:1이라 순서를 바꾸지 않는다.
# 적합도 점수·근거는 뒤에 덧붙인다 (evolve의 수확 코드가 인덱스로 읽으므로 앞을 건드리면 안 됨).
EXTRA_COLUMNS = ["적합도 점수", "평가 근거"]
WIDTHS = [10, 26, 40, 24, 20, 12, 26, 22, 26, 40, 22, 10, 14, 10, 34, 12, 60]

# 검수 결과 → 엑셀 '공지 상태'
DECISION_LABEL = {"확정": "확정", "기각": "기각", "보류": "보류"}


def _row(c) -> list:
    left = db.days_left(c["deadline"])
    # 사람이 판단했으면 그 값, 아니면 파이프라인과 같은 표기를 쓴다
    # (auto_status는 내부값이라 그대로 내보내면 'LLM판정_부적합' 같은 게 엑셀에 뜬다)
    status = (DECISION_LABEL.get(c["decision"])
              or STATUS_LABEL.get(c["auto_status"], "신규후보(확인필요)"))
    remark = (f"{c['source']} 수집 / 판정: 재학생 {c['judge_student'] or '?'}, "
              f"팀 {c['judge_team'] or '?'}, D-{left if left is not None else '?'}")
    if c["decision"] == "기각" and c["reason"]:
        remark += f" / 기각 사유: {db.REJECT_REASONS[c['reason']]['label']}"
    if c["note"]:
        remark += f" / {c['note']}"
    return [c["first_seen"][:7], c["school"] or "미분류", c["title"], c["host"] or "",
            c["field"] or "", c["deadline"] or "확인필요", c["schedule"] or "확인필요",
            c["team_size"] or "확인필요", c["eligibility"] or "확인필요", c["link"] or "",
            c["prize"] or "", c["fee"] or "확인필요", status, c["owner"] or "", remark]


def export_xlsx(target, only: str = "") -> int:
    """DB → 공지용 엑셀. only='확정'이면 확정 건만. → 행 수."""
    rows = db.list_candidates(decision=only) if only else db.list_candidates()
    wb = Workbook()
    ws = wb.active
    ws.title = "후보목록"
    ws.append(COLUMNS + EXTRA_COLUMNS)
    for cell, w in zip(ws[1], WIDTHS):
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.column_dimensions[cell.column_letter].width = w
    ws.freeze_panes = "A2"
    for c in rows:
        pts, gr, why = scoring.summary(c)
        ws.append(_row(c) + [f"{pts} ({gr})", why])
        for cell in ws[ws.max_row]:
            cell.font = BODY_FONT
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    wb.save(target)
    return len(rows)


if __name__ == "__main__":
    db.init()
    n = db.import_from_xlsx(os.path.join(
        os.path.dirname(os.path.dirname(os.path.abspath(__file__))),
        "data", "후보_컴피티션_목록.xlsx"))
    print(f"엑셀에서 {n}건 적재. 현재: {db.counts()}")
